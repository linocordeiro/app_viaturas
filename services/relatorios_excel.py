import io
from collections.abc import Iterable
from typing import TYPE_CHECKING, Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from fichas.models import FichaControle
    from veiculos.models import Manutencao, Viatura

# Estilos Frontline PF para Excel
HEADER_FILL = PatternFill(start_color="111213", end_color="111213", fill_type="solid")
GOLD_FILL = PatternFill(start_color="E1AD62", end_color="E1AD62", fill_type="solid")
ZEBRA_FILL = PatternFill(start_color="F4F6F8", end_color="F4F6F8", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="111213")
SUBTITLE_FONT = Font(name="Calibri", size=10, bold=False, color="555868")
BOLD_FONT = Font(name="Calibri", size=10, bold=True, color="111213")
REGULAR_FONT = Font(name="Calibri", size=10, color="41434E")

THIN_BORDER = Border(
    left=Side(style="thin", color="DADADD"),
    right=Side(style="thin", color="DADADD"),
    top=Side(style="thin", color="DADADD"),
    bottom=Side(style="thin", color="DADADD"),
)


def auto_ajustar_colunas(ws: Any) -> None:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def gerar_excel_ficha(ficha: "FichaControle") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active

    ws.title = f"Ficha {ficha.data_expediente.strftime('%d-%m-%Y')}"

    # Cabeçalho Institucional
    ws.merge_cells("A1:G1")
    ws["A1"] = "POLÍCIA FEDERAL — DEPARTAMENTO DE POLÍCIA FEDERAL"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = f"FICHA DIÁRIA DE CONTROLE DE VIATURAS — EXPEDIENTE {ficha.data_expediente.strftime('%d/%m/%Y')}"
    ws["A2"].font = Font(name="Calibri", size=12, bold=True, color="3363CC")
    ws["A2"].alignment = Alignment(horizontal="center")

    # Informações do Expediente
    ws["A4"] = "Data do Expediente:"
    ws["B4"] = ficha.data_expediente.strftime("%d/%m/%Y")
    ws["C4"] = "Horário de Início:"
    ws["D4"] = ficha.horario_inicio.strftime("%H:%M")
    ws["E4"] = "Horário de Término:"
    ws["F4"] = ficha.horario_termino.strftime("%H:%M")

    ws["A5"] = "Vigilante do Dia:"
    ws["B5"] = ficha.nome_vigilante
    ws["C5"] = "Status da Ficha:"
    ws["D5"] = ficha.get_status_display()
    ws["E5"] = "Visto Responsável:"
    ws["F5"] = "ASSINADO" if ficha.visto_responsavel else "Pendente"
    ws["G5"] = f"Visto Chefia: {'ASSINADO' if ficha.visto_chefia else 'Pendente'}"

    for r in range(4, 6):
        for c in range(1, 8):
            cell = ws.cell(row=r, column=c)
            cell.font = BOLD_FONT if c % 2 != 0 else REGULAR_FONT
            cell.fill = ZEBRA_FILL
            cell.border = THIN_BORDER

    # Tabela de Movimentações
    headers = [
        "Viatura / Modelo", "Placa", "Condutor", "Destino / Missão",
        "Saída (Hora/KM)", "Chegada (Hora/KM)", "KM Percorrido", "Possui Avarias?", "Avarias Encontradas"
    ]

    start_row = 7
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    current_row = start_row + 1
    for reg in ficha.registros.all():
        saida_str = f"{reg.horario_saida.strftime('%H:%M')} ({reg.odometro_saida:,} km)"
        chegada_str = f"{reg.horario_chegada.strftime('%H:%M')} ({reg.odometro_chegada:,} km)" if reg.horario_chegada and reg.odometro_chegada else "EM TRÂNSITO"
        km_str = reg.km_percorrido if reg.odometro_chegada else 0
        avaria_str = "SIM" if reg.possui_avarias else "NÃO"

        row_data = [
            f"{reg.viatura.marca} {reg.viatura.modelo}",
            reg.viatura.placa,
            reg.condutor,
            reg.destino,
            saida_str,
            chegada_str,
            km_str,
            avaria_str,
            reg.avarias_encontradas or "-"
        ]

        fill = ZEBRA_FILL if current_row % 2 == 0 else WHITE_FILL
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = REGULAR_FONT
            cell.fill = fill
            cell.border = THIN_BORDER
            if col_idx in [2, 5, 6, 7, 8]:
                cell.alignment = Alignment(horizontal="center")
        current_row += 1

    auto_ajustar_colunas(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def gerar_excel_viaturas(viaturas: Iterable["Viatura"]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Frota de Viaturas"

    ws.merge_cells("A1:H1")
    ws["A1"] = "POLÍCIA FEDERAL — RELATÓRIO GERAL DA FROTA DE VEÍCULOS"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = [
        "Placa", "Marca / Modelo", "Ano/Modelo", "Tipo", "Setor Pertencente",
        "Responsável", "KM Atual", "Próxima Manutenção", "Estado"
    ]

    start_row = 3
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER

    current_row = start_row + 1
    for v in viaturas:
        proxima = []
        if v.proxima_manutencao_km:
            proxima.append(f"{v.proxima_manutencao_km:,} km")
        if v.proxima_manutencao_data:
            proxima.append(v.proxima_manutencao_data.strftime("%d/%m/%Y"))
        proxima_str = " / ".join(proxima) if proxima else "Não agendada"

        row_data = [
            v.placa,
            f"{v.marca} {v.modelo}",
            f"{v.ano_fabricacao}/{v.ano_modelo}",
            v.get_tipo_display(),
            v.setor_pertencente.sigla,
            v.identificacao_responsavel,
            v.km_atual,
            proxima_str,
            v.get_status_display()
        ]

        fill = ZEBRA_FILL if current_row % 2 == 0 else WHITE_FILL
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = REGULAR_FONT
            cell.fill = fill
            cell.border = THIN_BORDER
            if col_idx in [1, 3, 7, 9]:
                cell.alignment = Alignment(horizontal="center")
        current_row += 1

    auto_ajustar_colunas(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def gerar_excel_manutencoes(
    viatura: "Viatura",
    manutencoes: Iterable["Manutencao"],
) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Manutenções {viatura.placa}"

    ws.merge_cells("A1:G1")
    ws["A1"] = f"POLÍCIA FEDERAL — HISTÓRICO DE MANUTENÇÃO: {viatura.marca} {viatura.modelo} ({viatura.placa})"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = [
        "Data do Serviço", "Tipo", "Odômetro (KM)", "Fornecedor / Oficina",
        "Nº OS / NF", "Descrição dos Serviços Realizados", "Peças Substituídas", "Valor (R$)"
    ]

    start_row = 3
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER

    current_row = start_row + 1
    total = 0
    for m in manutencoes:
        total += float(m.valor_total)
        row_data = [
            m.data_manutencao.strftime("%d/%m/%Y"),
            m.get_tipo_display(),
            m.km_no_momento,
            m.fornecedor_oficina,
            m.numero_ordem_servico or "-",
            m.descricao_servico,
            m.pecas_substituidas or "-",
            float(m.valor_total)
        ]

        fill = ZEBRA_FILL if current_row % 2 == 0 else WHITE_FILL
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = REGULAR_FONT
            cell.fill = fill
            cell.border = THIN_BORDER
            if col_idx in [1, 3]:
                cell.alignment = Alignment(horizontal="center")
            if col_idx == 8:
                cell.number_format = "R$ #,##0.00"
        current_row += 1

    # Linha de Total
    ws.cell(row=current_row, column=7, value="TOTAL INVESTIDO:").font = BOLD_FONT
    total_cell = ws.cell(row=current_row, column=8, value=total)
    total_cell.font = BOLD_FONT
    total_cell.number_format = "R$ #,##0.00"
    total_cell.fill = GOLD_FILL

    auto_ajustar_colunas(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
